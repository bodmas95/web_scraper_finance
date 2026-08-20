import openpyxl
import io
from typing import Any


async def parse_bref_template(excel_bytes: bytes) -> dict[str, Any]:
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    result = {}

    sheet_mapping = {
        "Input - Income Statement": "income_statement",
        "Input - Assets": "balance_sheet",
        "Input - Liabilities": "balance_sheet",
        "Input - Cash flow": "cash_flow",
        "Input - Cash Flow": "cash_flow",
    }

    for sheet_name, key in sheet_mapping.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        fields = {}
        for row in ws.iter_rows(min_row=3, values_only=False):
            field_name = row[0].value
            if not field_name or str(field_name).strip().startswith("INPUT"):
                continue
            values = {}
            for col_idx, header in enumerate(headers):
                if header and col_idx > 0 and header != "Var":
                    cell_val = row[col_idx].value
                    if cell_val is not None and not isinstance(cell_val, str):
                        year = str(header).split("\n")[0].split("/")[-1] if "/" in str(header) else str(header)
                        values[year.strip()] = cell_val
            if values:
                fields[field_name.strip()] = values

        if key in result:
            result[key].update(fields)
        else:
            result[key] = fields

    wb.close()
    return result


async def export_bref_to_excel(mapping_data: dict) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BREF Summary"

    ws.append(["Field", "Previous Year", "Current Year", "Match Method", "Confidence", "Source"])

    for statement_type in ["income_statement", "balance_sheet", "cash_flow"]:
        statement = mapping_data.get(statement_type, {})
        if not statement:
            continue
        ws.append([statement_type.upper().replace("_", " ")])
        for field_name, data in statement.items():
            if isinstance(data, dict):
                ws.append([
                    field_name,
                    data.get("previous_year", ""),
                    data.get("current_year", ""),
                    data.get("match_method", ""),
                    data.get("confidence", ""),
                    data.get("source_field", ""),
                ])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
